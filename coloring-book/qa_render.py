"""Render each slide from the built .pptx so it can be inspected.

LibreOffice cannot convert in this sandbox, so this reads the real shape
geometry back out of the file with python-pptx and redraws it with PIL. Text
is measured with DejaVu, which is wider than Calibri/Cambria, so anything that
fits here fits in PowerPoint.
"""
import glob
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

DPI = 96.0
PX = DPI / 914400.0            # EMU -> px
PT = DPI / 72.0                # pt  -> px
SERIF = "/usr/share/fonts/truetype/dejavu/DejaVuSerif%s.ttf"
SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf"


def font_for(name, size_pt, bold, italic):
    base = SERIF if (name or "").lower().startswith(("cambria", "bookman", "century")) else SANS
    suffix = "-Bold" if bold else ("-Oblique" if italic else "")
    try:
        return ImageFont.truetype(base % suffix, max(int(size_pt * PT), 6))
    except OSError:
        return ImageFont.truetype(base % "", max(int(size_pt * PT), 6))


def wrap(draw, text, font, max_w):
    lines = []
    for para in text.split("\n"):
        words, cur = para.split(), ""
        for w in words:
            t = (cur + " " + w).strip()
            if draw.textlength(t, font=font) <= max_w or not cur:
                cur = t
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def render(pptx_path, outdir):
    prs = Presentation(pptx_path)
    W = int(prs.slide_width * PX)
    H = int(prs.slide_height * PX)
    problems = []
    for idx, slide in enumerate(prs.slides, 1):
        bg = "#FFFFFF"
        try:
            if slide.background.fill.type is not None and slide.background.fill.type == 1:
                bg = "#" + str(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        im = Image.new("RGB", (W, H), bg)
        d = ImageDraw.Draw(im)
        boxes = []
        for sh in slide.shapes:
            x, y = sh.left * PX, sh.top * PX
            w, h = sh.width * PX, sh.height * PX
            if x < -1 or y < -1 or x + w > W + 1 or y + h > H + 1:
                problems.append(f"slide {idx}: shape out of bounds ({sh.shape_type})")
            if sh.has_text_frame and sh.text_frame.text.strip():
                boxes.append((x, y, w, h, sh.text_frame.text.strip()[:26]))
            if sh.shape_type == 13:                       # picture
                blob = sh.image.blob
                tmp = os.path.join(outdir, "_tmp_img")
                with open(tmp, "wb") as fh:
                    fh.write(blob)
                pic = Image.open(tmp).convert("RGB").resize((max(int(w), 1), max(int(h), 1)))
                im.paste(pic, (int(x), int(y)))
                continue
            fill = None
            try:
                if sh.fill.type == 1:
                    fill = "#" + str(sh.fill.fore_color.rgb)
            except Exception:
                pass
            if fill:
                if "ellipse" in str(sh.shape_type).lower() or sh.shape_type == 9:
                    d.ellipse([x, y, x + w, y + h], fill=fill)
                else:
                    d.rounded_rectangle([x, y, x + w, y + h], radius=8, fill=fill)
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            vcenter = str(getattr(sh.text_frame, "vertical_anchor", "") or "")
            tf = sh.text_frame
            ty = y + 2
            used = 0.0
            first_txt = tf.text.strip().replace("\n", " ")[:38]
            for para in tf.paragraphs:
                txt = "".join(r.text for r in para.runs)
                if not txt.strip():
                    continue
                r0 = para.runs[0]
                size = (r0.font.size.pt if r0.font.size else 18)
                col = "#1F2233"
                try:
                    if r0.font.color and r0.font.color.rgb:
                        col = "#" + str(r0.font.color.rgb)
                except Exception:
                    pass
                fo = font_for(r0.font.name, size, bool(r0.font.bold), bool(r0.font.italic))
                bullet = para.level >= 0 and "buChar" in para._pPr.xml if para._pPr is not None else False
                prefix = "•  " if bullet else ""
                lines = wrap(d, prefix + txt, fo, max(w - 6, 10))
                lh = size * PT * 1.21
                total = len(lines) * lh
                align = str(para.alignment or "")
                used += total + (size * PT * 0.42 if used else 0)
                for ln in lines:
                    tw = d.textlength(ln, font=fo)
                    tx = x + 3
                    if "CENTER" in align:
                        tx = x + (w - tw) / 2
                    d.text((tx, ty), ln, font=fo, fill=col)
                    ty += lh
                ty += size * PT * 0.42
            if used - 6 > h:
                problems.append(f"slide {idx}: text overflows ({first_txt!r} "
                                f"needs {used/PT:.0f}pt in {h/PT:.0f}pt)")
        im.save(os.path.join(outdir, f"slide-{idx:02d}.jpg"), quality=88)
    tmp = os.path.join(outdir, "_tmp_img")
    if os.path.exists(tmp):
        os.remove(tmp)
    return len(prs.slides), problems


if __name__ == "__main__":
    outdir = sys.argv[2] if len(sys.argv) > 2 else "qa"
    os.makedirs(outdir, exist_ok=True)
    for f in glob.glob(os.path.join(outdir, "*.jpg")):
        os.remove(f)
    n, probs = render(sys.argv[1], outdir)
    print(f"rendered {n} slides -> {outdir}/")
    if probs:
        print("\npotential problems:")
        for p in probs:
            print("  " + p)
    else:
        print("no bounds or overflow problems detected")
